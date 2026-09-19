"""Format eksportu/importu panelu sklepu.

Plik z panelu (`admin-product-product-*.xlsx`) ma stały układ:

    wiersz 1   hash eksportu (kolumna 27)
    wiersz 2   „Produkty - Produkty"
    wiersz 3   opis filtrów, z jakimi zrobiono eksport
    wiersz 4   Export date | data | Items: | N | Limit: 2000
    wiersz 5   pusty
    wiersz 6   NAGŁÓWKI (165 kolumn)
    wiersz 7+  dane

Najważniejsza rzecz w tym module: wartości słownikowe w panelu to
`ID|etykieta` („2022|tapicerowane"), a nasze źródło (eksport z BigQuery) ma
same etykiety. Bez ID plik się nie zaimportuje albo — gorzej — zaimportuje
się jako nowa wartość słownika. Dlatego mapowanie etykieta→ID wyciągamy
z prawdziwych plików z panelu i trzymamy w `config/slownik_idow.yaml`.

Konsekwencja, o której trzeba pamiętać: słownik zna tylko te wartości,
które kiedykolwiek widzieliśmy w jakimś eksporcie. Poprawki na wartości
spoza słownika nie trafiają do pliku — lądują na liście „bez ID" na
stronie eksportu, zamiast po cichu wyjść jako gołe etykiety.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import yaml

from .config import KATALOG_CONFIG
from .tekst import norm

PLIK_SLOWNIKA = KATALOG_CONFIG / "slownik_idow.yaml"
PLIK_WZORCA = KATALOG_CONFIG / "wzorzec_panelu.yaml"

WIERSZ_NAGLOWKA = 6
RE_ID_ETYKIETA = re.compile(r"^(\d+)\|(.+)$", re.S)

# Kolumny, po których panel rozpoznaje produkt — zawsze wypełniane.
KOLUMNY_KLUCZA = ("ID", "Kod", "Kod producenta", "Nazwa")

# Nasze nazwy atrybutów = nazwy kolumn w panelu, poza tymi wyjątkami.
ALIASY_KOLUMN: dict[str, str] = {
    "Na nóżkach": "Na nóżkach",
    "Liczba drzwi": "Liczba drzwi",
}


@dataclass
class Wzorzec:
    """Układ pliku z panelu — nagłówki, preambuła i charakter kolumn."""
    naglowki: list[str] = field(default_factory=list)
    preambula: list[list] = field(default_factory=list)   # wiersze 1-5
    # kolumny, w których widzieliśmy wartości surowe (liczby, teksty) — te
    # wolno przepisać jeden do jednego
    surowe: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return bool(self.naglowki) and "ID" in self.naglowki


def wczytaj_wzorzec() -> Wzorzec:
    """Układ pliku panelu, z nazwami kolumn doprowadzonymi do dzisiejszych.

    Nagłówki pochodzą z eksportu produktów, więc zamarzają na dzień, w którym
    ten eksport powstał. Po przemianowaniu atrybutu w panelu zostawała stara
    nazwa kolumny („Ilość osób"), a słownik znał już nową („Liczba miejsc") —
    i atrybut po cichu wypadał z pliku importu jako „kolumny nie ma
    w formacie panelu". Mapa przemianowań jest tu stosowana przy każdym
    odczycie, więc naprawia się samo, nawet bez ponownego wgrania eksportu.
    """
    if not PLIK_WZORCA.exists():
        return Wzorzec()
    d = yaml.safe_load(PLIK_WZORCA.read_text(encoding="utf-8")) or {}
    mapa = wczytaj_zmiany_nazw()["atrybuty"]
    pod_nowa = (lambda x: mapa.get(x, x)) if mapa else (lambda x: x)
    return Wzorzec(naglowki=[pod_nowa(k) for k in (d.get("naglowki") or [])],
                   preambula=[list(w) for w in (d.get("preambula") or [])],
                   surowe=[pod_nowa(k) for k in (d.get("surowe") or [])])


def zapisz_wzorzec(w: Wzorzec) -> None:
    PLIK_WZORCA.parent.mkdir(parents=True, exist_ok=True)
    PLIK_WZORCA.write_text(yaml.safe_dump(
        {"naglowki": w.naglowki, "surowe": w.surowe, "preambula": w.preambula},
        allow_unicode=True, sort_keys=False), encoding="utf-8")


# --- słownik ID -----------------------------------------------------------

def wczytaj_slownik() -> dict[str, dict[str, list[str]]]:
    if not PLIK_SLOWNIKA.exists():
        return {}
    return yaml.safe_load(PLIK_SLOWNIKA.read_text(encoding="utf-8")) or {}


def zapisz_slownik(s: dict) -> None:
    PLIK_SLOWNIKA.parent.mkdir(parents=True, exist_ok=True)
    PLIK_SLOWNIKA.write_text(yaml.safe_dump(s, allow_unicode=True, sort_keys=True),
                             encoding="utf-8")


PLIK_ETYKIET = KATALOG_CONFIG / "etykiety_panelu.yaml"


def wczytaj_etykiety() -> dict[str, dict[str, str]]:
    if not PLIK_ETYKIET.exists():
        return {}
    return yaml.safe_load(PLIK_ETYKIET.read_text(encoding="utf-8")) or {}


def zapisz_etykiety(e: dict) -> None:
    PLIK_ETYKIET.parent.mkdir(parents=True, exist_ok=True)
    PLIK_ETYKIET.write_text(yaml.safe_dump(e, allow_unicode=True, sort_keys=True),
                            encoding="utf-8")


def id_dla(slownik: dict, kolumna: str, etykieta: str,
           etykiety: dict | None = None) -> tuple[str | None, str]:
    """(wartość do pliku, powód gdy się nie da).

    Zwraca `ID|etykieta`, gdy znamy dokładnie jedno ID. Gdy etykieta ma
    w panelu kilka ID (zdarza się — ta sama nazwa w różnych zestawach
    atrybutów), nie zgadujemy: lepiej zostawić to człowiekowi.
    """
    klucz = norm(etykieta)
    ids = (slownik.get(kolumna) or {}).get(klucz)
    if not ids:
        return None, "wartość spoza słownika sklepu"
    if len(ids) > 1:
        return None, f"niejednoznaczne ID ({', '.join(ids)})"
    kanoniczna = ((etykiety or {}).get(kolumna) or {}).get(klucz, etykieta)
    return f"{ids[0]}|{kanoniczna}", ""


# Wartości wielokrotne: nasz eksport podaje je jako „ceramika, metal".
# Każdy człon ma własne ID, więc mapujemy je po kolei i sklejamy z powrotem.
ROZDZIELACZ = ", "


def wartosc_slownikowa(slownik: dict, kolumna: str, wartosc: str,
                       etykiety: dict | None = None) -> tuple[str | None, str]:
    czlony = [c.strip() for c in str(wartosc).split(",") if c.strip()]
    if len(czlony) <= 1:
        return id_dla(slownik, kolumna, wartosc, etykiety)

    wyniki = []
    for c in czlony:
        w, powod = id_dla(slownik, kolumna, c, etykiety)
        if powod:
            return None, f"{powod}: „{c}”"
        wyniki.append(w)
    return ROZDZIELACZ.join(wyniki), ""


def wartosc_do_pliku(slownik: dict, kolumna: str, wartosc: str,
                     kolumny_slownikowe: set[str],
                     kolumny_surowe: set[str] | None = None,
                     etykiety: dict | None = None) -> tuple[str | None, str]:
    """Liczby i teksty idą jak są; słownikowe muszą dostać ID.

    Kolumna, której nigdy nie widzieliśmy w pliku z panelu, jest trzecim
    przypadkiem i najgroźniejszym: nie wiemy, czy panel oczekuje tam gołego
    tekstu, czy `ID|etykieta`. Wpisanie gołej etykiety do kolumny słownikowej
    kończy się nową, śmieciową wartością w słowniku sklepu — więc takiej
    poprawki po prostu nie wypuszczamy.
    """
    if kolumna in kolumny_slownikowe:
        return wartosc_slownikowa(slownik, kolumna, wartosc, etykiety)
    if kolumny_surowe is None or kolumna in kolumny_surowe:
        return wartosc, ""
    return None, "kolumny nie było w żadnym pliku z panelu"


# --- nauka z pliku panelu -------------------------------------------------

def czy_plik_panelu(sciezka: str | Path) -> bool:
    """Czy to eksport z panelu (a nie nasz eksport produktów z BigQuery)."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(sciezka, read_only=True)
        ws = wb.active
        naglowki = [c.value for c in next(ws.iter_rows(
            min_row=WIERSZ_NAGLOWKA, max_row=WIERSZ_NAGLOWKA))]
        wb.close()
        return "ID" in naglowki and "Kod producenta" in naglowki
    except Exception:
        return False


def naucz_z_pliku(sciezka: str | Path) -> dict:
    """Wyciąga z pliku panelu układ kolumn i wszystkie pary ID|etykieta.

    Scala ze stanem na dysku — każdy kolejny eksport z panelu dokłada
    wartości, których wcześniej nie widzieliśmy.
    """
    import openpyxl
    wb = openpyxl.load_workbook(sciezka, read_only=True)
    ws = wb.active

    wiersze = list(ws.iter_rows(values_only=True))
    wb.close()
    if len(wiersze) < WIERSZ_NAGLOWKA:
        raise ValueError("plik nie wygląda na eksport z panelu")

    naglowki = [(h or "") for h in wiersze[WIERSZ_NAGLOWKA - 1]]
    preambula = [[("" if c is None else c) for c in w]
                 for w in wiersze[:WIERSZ_NAGLOWKA - 1]]

    slownik = wczytaj_slownik()
    stary_wzorzec = wczytaj_wzorzec()
    nowych = 0
    kolumny_slownikowe: set[str] = set()
    surowe: set[str] = set(stary_wzorzec.surowe)

    for wiersz in wiersze[WIERSZ_NAGLOWKA:]:
        for naglowek, wartosc in zip(naglowki, wiersz):
            if not naglowek or wartosc is None or str(wartosc).strip() == "":
                continue
            m = RE_ID_ETYKIETA.match(str(wartosc).strip())
            if not m:
                surowe.add(naglowek)
                continue
            kolumny_slownikowe.add(naglowek)
            ident, etykieta = m.group(1), m.group(2).strip()
            kubelek = slownik.setdefault(naglowek, {})
            lista = kubelek.setdefault(norm(etykieta), [])
            if ident not in lista:
                lista.append(ident)
                nowych += 1

    zapisz_slownik(slownik)
    # kolumna, która kiedykolwiek wystąpiła jako ID|etykieta, zostaje
    # słownikowa — nawet jeśli gdzie indziej trafiła się w niej surowa wartość
    surowe -= set(slownik.keys())
    zapisz_wzorzec(Wzorzec(naglowki=naglowki, preambula=preambula,
                           surowe=sorted(surowe)))

    return {
        "kolumny": len(naglowki),
        "produktow": len(wiersze) - WIERSZ_NAGLOWKA,
        "kolumny_slownikowe": sorted(kolumny_slownikowe),
        "kolumny_surowe": sorted(surowe),
        "nowych_wartosci": nowych,
        "wartosci_razem": sum(len(v) for v in slownik.values()),
        "nierozpoznanych": len([h for h in naglowki
                                if h and h not in slownik and h not in surowe]),
    }


def kolumny_slownikowe(slownik: dict | None = None) -> set[str]:
    return set((slownik if slownik is not None else wczytaj_slownik()).keys())


# --- pełny słownik atrybutów ze sklepu ------------------------------------
#
# Eksport „Atrybuty" z panelu ma dwa arkusze: listę atrybutów (id, status,
# title) i listę wartości słownikowych (id, atrybut, wartość). To jest
# źródło prawdy, w przeciwieństwie do uczenia się z eksportów produktów,
# które pokazują tylko te wartości, które akurat gdzieś wystąpiły.

PLIK_ATRYBUTOW = KATALOG_CONFIG / "atrybuty_panelu.yaml"


NAGLOWEK_ATRYBUTOW = ("id", "status", "title")
NAGLOWEK_WARTOSCI = ("id", "title", "title")


def _naglowek(ws) -> tuple:
    try:
        return tuple((str(c).strip().lower() if c is not None else "")
                     for c in next(ws.iter_rows(values_only=True))[:3])
    except StopIteration:
        return ()


def pary_arkuszy(wb) -> list[tuple]:
    """Pary (arkusz atrybutów, arkusz wartości) w kolejności z pliku.

    Panel dokłada kolejne zrzuty słownika jako NOWE arkusze („atrybuty 9.18",
    „wartości 9.18") i zostawia stare obok. Branie dwóch pierwszych arkuszy
    uczyło nas wtedy po cichu nieaktualnego słownika — plik wygląda na wgrany,
    a zmiany z panelu nie wchodzą. Dlatego rozpoznajemy arkusze po nagłówku,
    a bierzemy ostatnią parę: zrzuty idą chronologicznie, najnowszy na końcu.
    """
    pary, otwarty = [], None
    for ws in wb.worksheets:
        n = _naglowek(ws)
        if n == NAGLOWEK_ATRYBUTOW:
            otwarty = ws
        elif n == NAGLOWEK_WARTOSCI and otwarty is not None:
            pary.append((otwarty, ws))
            otwarty = None
    return pary


def czy_plik_slownika(sciezka: str | Path) -> bool:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(sciezka, read_only=True)
        ile = len(pary_arkuszy(wb))
        wb.close()
        return ile > 0
    except Exception:
        return False


def _ident(v) -> str:
    """'2022.0' i 2022 to ten sam identyfikator — arkusz podaje je jako float."""
    try:
        return str(int(float(v)))
    except (TypeError, ValueError):
        return str(v).strip()


def wczytaj_atrybuty_panelu() -> dict[str, dict]:
    if not PLIK_ATRYBUTOW.exists():
        return {}
    return yaml.safe_load(PLIK_ATRYBUTOW.read_text(encoding="utf-8")) or {}


def naucz_ze_slownika(sciezka: str | Path) -> dict:
    """Wgrywa kompletny słownik: atrybuty + ich dozwolone wartości.

    Nadpisuje wartości poznane wcześniej z eksportów produktów — ten plik
    jest pełny, tamte były wyrywkowe.
    """
    import openpyxl
    wb = openpyxl.load_workbook(sciezka, read_only=True)
    pary = pary_arkuszy(wb)
    if not pary:
        wb.close()
        raise ValueError("w pliku nie ma arkuszy słownika "
                         "(nagłówki id/status/title oraz id/title/title)")
    # ostatnia para = najnowszy zrzut; poprzednie zostawiamy w spokoju
    ark_atrybuty, ark_wartosci = pary[-1]
    arkusze = f"{ark_atrybuty.title} + {ark_wartosci.title}"

    stare_atrybuty = wczytaj_atrybuty_panelu()
    stary_slownik, stare_etykiety = wczytaj_slownik(), wczytaj_etykiety()

    atrybuty: dict[str, dict] = {}
    for r in list(ark_atrybuty.iter_rows(values_only=True))[1:]:
        if r[0] is None or not r[2]:
            continue
        atrybuty[str(r[2]).strip()] = {"id": _ident(r[0]), "status": str(r[1] or "")}

    slownik: dict[str, dict[str, list[str]]] = {}
    etykiety: dict[str, dict[str, str]] = {}
    ile_wartosci = 0
    smieci = 0
    for r in list(ark_wartosci.iter_rows(values_only=True))[1:]:
        if r[0] is None or not r[1] or r[2] is None:
            continue
        kolumna, etykieta = str(r[1]).strip(), str(r[2]).strip()
        # w słowniku sklepu siedzą pozycje o tytule „NULL" — to śmieć po
        # imporcie, nigdy poprawny cel poprawki
        if not etykieta or etykieta.upper() == "NULL":
            smieci += 1
            continue
        lista = slownik.setdefault(kolumna, {}).setdefault(norm(etykieta), [])
        # do pliku ma iść pisownia sklepu, nie nasza z eksportu BigQuery
        etykiety.setdefault(kolumna, {})[norm(etykieta)] = etykieta
        ident = _ident(r[0])
        if ident not in lista:
            lista.append(ident)
            ile_wartosci += 1
    wb.close()

    zapisz_slownik(slownik)
    zapisz_etykiety(etykiety)
    PLIK_ATRYBUTOW.parent.mkdir(parents=True, exist_ok=True)
    PLIK_ATRYBUTOW.write_text(yaml.safe_dump(atrybuty, allow_unicode=True, sort_keys=True),
                              encoding="utf-8")

    # atrybut z listy, którego nie ma w wartościach, jest polem wolnym
    # (liczba albo tekst) — te wolno przepisywać jeden do jednego
    wzorzec = wczytaj_wzorzec()
    wzorzec.surowe = sorted(set(wzorzec.surowe) |
                            {k for k in atrybuty if k not in slownik})
    wzorzec.surowe = [k for k in wzorzec.surowe if k not in slownik]
    zapisz_wzorzec(wzorzec)

    zmiany = _wykryj_zmiany_nazw(stare_atrybuty, atrybuty,
                                 stary_slownik, stare_etykiety, slownik, etykiety)
    migrowane: list[str] = []
    if zmiany["atrybuty"] or zmiany["wartosci"]:
        dopisz_zmiany_nazw(zmiany)
        migrowane = migruj_config(zmiany)

    wylaczone = [k for k, v in atrybuty.items() if v["status"] != "ACTIVE"]
    return {
        "arkusze": arkusze,
        "przemianowanych_atrybutow": len(zmiany["atrybuty"]),
        "przemianowanych_wartosci": sum(len(v) for v in zmiany["wartosci"].values()),
        "zmiany_nazw": zmiany,
        "zmigrowane_pliki": migrowane,
        "atrybutow": len(atrybuty),
        "slownikowych": len(slownik),
        "wartosci": ile_wartosci,
        "wolnych": len([k for k in atrybuty if k not in slownik]),
        "wylaczonych": len(wylaczone),
        "wylaczone": sorted(wylaczone),
        "smieci": smieci,
        "duplikaty": duplikaty(slownik),
    }


# --- przemianowania w słowniku sklepu -------------------------------------
#
# Panel zmienia tytuł atrybutu albo wartości, a ID zostaje. Dla sklepu to
# kosmetyka, dla nas nie: eksporty produktów sprzed zmiany dalej niosą starą
# nazwę („Materiał obicia"), a reguły, schema i słowniki mówią już nową
# („Rodzaj obicia"). Bez mapy stara→nowa taki atrybut po cichu wypada
# z walidacji — nie ma go w definicjach, więc nikt go nie sprawdza.
#
# Mapę budujemy po ID, bo tylko ID przeżywa zmianę tytułu.

PLIK_ZMIAN_NAZW = KATALOG_CONFIG / "zmiany_nazw.yaml"


def wczytaj_zmiany_nazw() -> dict:
    if not PLIK_ZMIAN_NAZW.exists():
        return {"atrybuty": {}, "wartosci": {}}
    d = yaml.safe_load(PLIK_ZMIAN_NAZW.read_text(encoding="utf-8")) or {}
    return {"atrybuty": d.get("atrybuty") or {}, "wartosci": d.get("wartosci") or {}}


def dopisz_zmiany_nazw(nowe: dict) -> dict:
    """Dokłada do mapy, nie nadpisuje — historia przemianowań się kumuluje.

    Gdy „A" stało się „B", a potem „B" stało się „C", stary eksport z „A"
    ma trafić na „C", więc przy dopisaniu przepinamy też wcześniejsze wpisy.
    """
    mapa = wczytaj_zmiany_nazw()

    for stara, nowa in (nowe.get("atrybuty") or {}).items():
        for k, v in list(mapa["atrybuty"].items()):
            if v == stara:
                mapa["atrybuty"][k] = nowa
        if stara != nowa:
            mapa["atrybuty"][stara] = nowa

    for atrybut, pary in (nowe.get("wartosci") or {}).items():
        cel = mapa["wartosci"].setdefault(atrybut, {})
        for stara, nowa in pary.items():
            for k, v in list(cel.items()):
                if v == stara:
                    cel[k] = nowa
            if stara != nowa:
                cel[stara] = nowa

    # atrybut przemianowany drugi raz zabiera ze sobą swoje wartości
    for stara, nowa in mapa["atrybuty"].items():
        if stara in mapa["wartosci"] and stara != nowa:
            mapa["wartosci"].setdefault(nowa, {}).update(mapa["wartosci"].pop(stara))

    PLIK_ZMIAN_NAZW.parent.mkdir(parents=True, exist_ok=True)
    PLIK_ZMIAN_NAZW.write_text(
        "# Mapa stara nazwa -> nowa nazwa, budowana po ID przy wgrywaniu\n"
        "# słownika z panelu. Plik jest nadpisywany przez aplikację.\n"
        + yaml.safe_dump(mapa, allow_unicode=True, sort_keys=True),
        encoding="utf-8")
    return mapa


def migruj_config(zmiany: dict) -> list[str]:
    """Przepisuje nazwy atrybutów w plikach pisanych maszynowo.

    `schema.yaml` (zakresy i pokrycie per kategoria) i `reguly.yaml` odwołują
    się do atrybutów po nazwie. Po przemianowaniu w panelu zostałyby przy
    starej i przestałyby cokolwiek łapać. `slowniki.yaml` zostawiamy człowiekowi
    — jest pisany ręcznie, z komentarzami, których nie chcemy zgubić.
    """
    from . import config
    mapa = {k: v for k, v in (zmiany.get("atrybuty") or {}).items() if k != v}
    if not mapa:
        return []

    def przepisz(x):
        if isinstance(x, dict):
            return {mapa.get(k, k) if isinstance(k, str) else k: przepisz(v)
                    for k, v in x.items()}
        if isinstance(x, list):
            return [przepisz(v) for v in x]
        return mapa.get(x, x) if isinstance(x, str) else x

    ruszone = []
    sch = config.schema()
    if sch:
        nowy = przepisz(sch)
        if nowy != sch:
            config.zapisz_schema(nowy)
            ruszone.append("schema.yaml")
    reg = config.reguly()
    if reg:
        nowy = przepisz(reg)
        if nowy != reg:
            config.zapisz_reguly(nowy)
            ruszone.append("reguly.yaml")
    # Nagłówki pliku panelu też — inaczej atrybut wypada z importu jako
    # „kolumny nie ma w formacie panelu", a odczyt i tak je już podmienia.
    wz = wczytaj_wzorzec()
    surowy = yaml.safe_load(PLIK_WZORCA.read_text(encoding="utf-8")) or {} \
        if PLIK_WZORCA.exists() else {}
    if surowy and (list(surowy.get("naglowki") or []) != wz.naglowki
                   or list(surowy.get("surowe") or []) != wz.surowe):
        zapisz_wzorzec(wz)
        ruszone.append("wzorzec_panelu.yaml")
    return ruszone


def _po_id(slownik: dict, etykiety: dict) -> dict[str, tuple[str, str]]:
    """ID wartości -> (atrybut, etykieta). Tylko jednoznaczne ID."""
    ile: dict[str, int] = {}
    for m in slownik.values():
        for ids in m.values():
            for i in ids:
                ile[i] = ile.get(i, 0) + 1
    out = {}
    for kolumna, m in slownik.items():
        for klucz, ids in m.items():
            if len(ids) != 1 or ile.get(ids[0], 0) != 1:
                continue
            out[ids[0]] = (kolumna, (etykiety.get(kolumna) or {}).get(klucz, klucz))
    return out


def _wykryj_zmiany_nazw(stare_atrybuty: dict, nowe_atrybuty: dict,
                        stary_slownik: dict, stare_etykiety: dict,
                        nowy_slownik: dict, nowe_etykiety: dict) -> dict:
    zmiany_atr: dict[str, str] = {}
    po_id = {v["id"]: k for k, v in stare_atrybuty.items() if v.get("id")}
    for nazwa, v in nowe_atrybuty.items():
        stara = po_id.get(v.get("id"))
        if stara and stara != nazwa:
            zmiany_atr[stara] = nazwa

    zmiany_wart: dict[str, dict[str, str]] = {}
    stare_po_id = _po_id(stary_slownik, stare_etykiety)
    nowe_po_id = _po_id(nowy_slownik, nowe_etykiety)
    for ident, (kolumna, etykieta) in nowe_po_id.items():
        poprzednie = stare_po_id.get(ident)
        if not poprzednie:
            continue
        stara_kolumna, stara_etykieta = poprzednie
        # zmiana nazwy atrybutu jest już zapisana wyżej — tu tylko wartości
        if zmiany_atr.get(stara_kolumna, stara_kolumna) != kolumna:
            continue
        if norm(stara_etykieta) != norm(etykieta):
            zmiany_wart.setdefault(kolumna, {})[stara_etykieta] = etykieta

    return {"atrybuty": zmiany_atr, "wartosci": zmiany_wart}


def duplikaty(slownik: dict | None = None) -> list[dict]:
    """Ta sama etykieta pod kilkoma ID w jednym atrybucie.

    To problem po stronie sklepu, nie nasz: dopóki „nowoczesny" ma trzy ID,
    nie da się rozstrzygnąć, które wpisać, więc takie poprawki wypadają
    z eksportu. Lista jest po to, żeby dało się to posprzątać w panelu.
    """
    s = slownik if slownik is not None else wczytaj_slownik()
    etykiety = wczytaj_etykiety()
    out = [{"atrybut": k, "wartosc": (etykiety.get(k) or {}).get(v, v), "ids": ids}
           for k, m in s.items() for v, ids in m.items() if len(ids) > 1]
    return sorted(out, key=lambda w: (w["atrybut"], w["wartosc"]))
